"""Turn report data into downloadable files: CSV, Excel (.xlsx) and the lender-ready PDF summary.

Safety: text from imported statements can start with "=", "+", "-" or "@", which a spreadsheet would run as a
formula (CSV/formula injection). Such cells are written with a leading apostrophe in CSV and as plain text in
Excel. Account numbers never appear beyond their last 4 digits.
"""
from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import fx

FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


@dataclass
class Column:
    key: str
    label: str
    kind: str = "text"  # text | money | number | percent | date


@dataclass
class Table:
    title: str
    columns: list[Column]
    rows: list[dict]
    meta: list[str] = field(default_factory=list)   # lines shown above the table in Excel
    styles: list[str] = field(default_factory=list)  # per row: "" | "header" | "subtotal" | "total"
    sheet: str = "Report"


def _safe_text(value) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(FORMULA_START) else text


def to_csv(table: Table) -> bytes:
    """UTF-8 with a byte-order mark, so Excel opens non-ASCII text correctly."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c.label for c in table.columns])
    for row in table.rows:
        w.writerow([_safe_text(row.get(c.key)) if c.kind == "text" else ("" if row.get(c.key) is None else row.get(c.key))
                    for c in table.columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(table: Table) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = table.sheet[:31]
    ws.append([table.title])
    ws["A1"].font = Font(bold=True, size=14)
    for line in table.meta:
        ws.append([line])
        ws.cell(ws.max_row, 1).font = Font(color="666666")
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.label for c in table.columns])
    thin = Side(style="thin", color="D0D3DE")
    for i in range(1, len(table.columns) + 1):
        cell = ws.cell(header_row, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4A22A3")
        cell.alignment = Alignment(horizontal="left" if table.columns[i - 1].kind == "text" else "right")
    formats = {"money": "#,##0.00;[Red]-#,##0.00", "number": "#,##0", "percent": "0.0%", "date": "yyyy-mm-dd"}
    for n, row in enumerate(table.rows):
        style = table.styles[n] if n < len(table.styles) else ""
        values = []
        for c in table.columns:
            v = row.get(c.key)
            if c.kind == "date" and isinstance(v, str) and v:
                v = date.fromisoformat(v)
            values.append(v)
        ws.append(values)
        r = ws.max_row
        for i, c in enumerate(table.columns, start=1):
            cell = ws.cell(r, i)
            if c.kind == "text":
                cell.value = "" if cell.value is None else str(cell.value)
                cell.data_type = "s"  # after the value (assigning re-detects "=..." as a formula): always plain text
            elif c.kind in formats:
                cell.number_format = formats[c.kind]
            if style in ("subtotal", "total", "header"):
                cell.font = Font(bold=True)
            if style == "total":
                cell.border = Border(top=thin, bottom=thin)
                cell.fill = PatternFill("solid", fgColor="F1EEFC")
    ws.freeze_panes = ws.cell(header_row + 1, 2 if table.columns and table.columns[0].kind == "text" else 1)
    for i, c in enumerate(table.columns, start=1):
        longest = max([len(c.label)] + [len(f"{row.get(c.key, '')}") for row in table.rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- report tables

def pnl_table(report: dict, meta: list[str], currency: str) -> Table:
    cols = [Column("label", "Line item")] + [Column(f"c{i}", f"{name} ({currency})", "money")
                                             for i, name in enumerate(report["columns"])]
    rows, styles = [], []
    for line in report["rows"]:
        row = {"label": line["label"]}
        if line["values"] is not None:
            row.update({f"c{i}": v for i, v in enumerate(line["values"])})
        rows.append(row)
        styles.append("" if line["kind"] == "item" else line["kind"])
    return Table("Profit and loss", cols, rows, meta, styles, sheet="Profit and loss")


def receivables_table(report: dict, meta: list[str], currency: str) -> Table:
    cols = [Column("customer", "Customer"), Column("invoices", "Open invoices", "number")] + \
           [Column(k, f"{'Current' if k == 'current' else k + ' days'} ({currency})", "money")
            for k in ("current", "1-30", "31-60", "61-90", "90+")] + \
           [Column("total", f"Total ({currency})", "money"), Column("oldestDaysPastDue", "Oldest (days past due)", "number")]
    rows = list(report["customers"]) + [{"customer": "Total", "invoices": sum(c["invoices"] for c in report["customers"]),
                                         **report["totals"]}]
    return Table("Aged receivables", cols, rows, meta, [""] * len(report["customers"]) + ["total"],
                 sheet="Aged receivables")


def forecast_table(report: dict, meta: list[str], currency: str) -> Table:
    m = lambda k, label: Column(k, f"{label} ({currency})", "money")  # noqa: E731
    cols = [Column("week", "Week", "number"), Column("start", "Week starting", "date"), m("opening", "Opening"),
            m("inflows", "Money in"), m("outflows", "Money out"), m("net", "Net"), m("closing", "Closing"),
            m("p10", "Cautious (P10)"), m("p90", "Optimistic (P90)"), Column("flag", "Note")]
    rows = [{**w, "flag": "Below safety buffer" if w["belowBuffer"] else ""} for w in report["weeks"]]
    return Table("13-week cash-flow forecast", cols, rows, meta, sheet="Cash-flow forecast")


def transactions_table(report: dict, meta: list[str], currency: str) -> Table:
    cols = [Column("date", "Date", "date"), Column("account", "Account"), Column("description", "Description"),
            Column("counterparty", "Counterparty"), Column("category", "Category"), Column("direction", "Direction"),
            Column("signed", f"Amount ({currency})", "money")]
    return Table("Transactions", cols, report["rows"], meta, sheet="Transactions")


# ---------------------------------------------------------------- lender summary PDF

def _latin(text) -> str:
    """The PDF uses the built-in Helvetica font (Latin-1): keep accents where possible, drop the rest."""
    text = "" if text is None else str(text)
    text = text.replace("—", "-").replace("–", "-").replace("’", "'").replace("“", '"').replace("”", '"').replace("…", "...")
    try:
        text.encode("latin-1")
        return text
    except UnicodeEncodeError:
        return unicodedata.normalize("NFKD", text).encode("latin-1", "ignore").decode("latin-1")


def summary_pdf(data: dict) -> bytes:
    """A two-page lender-ready financial summary. `data` is assembled by the reports API from the engines."""
    from fpdf import FPDF

    cur = data["currency"]
    m = lambda x: _latin(fx.fmt(x, cur))  # noqa: E731
    INK, MUTED, BRAND, LINE, SOFT = (20, 23, 43), (95, 102, 125), (74, 34, 163), (227, 230, 239), (243, 241, 252)

    class PDF(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", size=8)
            self.set_text_color(*MUTED)
            self.cell(0, 5, _latin(f"SME Finance Copilot  |  {data['business']['name']}  |  "
                                   f"Generated from the business's own records; not audited; not financial advice.  |  "
                                   f"Page {self.page_no()}/{{nb}}"), align="C")

    pdf = PDF(format="A4", unit="mm")
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(True, margin=16)
    pdf.set_margins(15, 14, 15)
    pdf.add_page()
    W = pdf.w - 30

    def heading(text: str):
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(*BRAND)
        pdf.cell(0, 7, _latin(text), new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(*LINE)
        pdf.line(15, pdf.get_y(), 15 + W, pdf.get_y())
        pdf.ln(2)
        pdf.set_text_color(*INK)

    def fit(text: str, width: float) -> str:
        """Shorten text to the column width at a word boundary, ending with '...'."""
        text = _latin(text)
        if pdf.get_string_width(text) <= width - 2:
            return text
        words = text.split()
        while words and pdf.get_string_width(" ".join(words) + "...") > width - 2:
            words.pop()
        return " ".join(words).rstrip(",;:(") + "..."

    def table(headers: list[str], rows: list[list[str]], widths: list[float], bold_last: bool = False,
              text_cols: tuple[int, ...] = (0,)):
        align = lambda i: "L" if i in text_cols else "R"  # noqa: E731  text left, figures right
        pdf.set_font("Helvetica", "B", 8.5)
        pdf.set_fill_color(*SOFT)
        gap = lambda i: "   " if i > 0 and i in text_cols else ""  # noqa: E731  space before a text column
        for i, (h, w) in enumerate(zip(headers, widths)):
            pdf.cell(w, 6, gap(i) + _latin(h), border="B", fill=True, align=align(i))
        pdf.ln()
        for n, row in enumerate(rows):
            last = bold_last and n == len(rows) - 1
            pdf.set_font("Helvetica", "B" if last else "", 8.5)
            for i, (v, w) in enumerate(zip(row, widths)):
                pdf.cell(w, 5.6, gap(i) + fit(v, w - 3 if gap(i) else w), border="T" if last else 0, align=align(i))
            pdf.ln()

    # ---- header
    b = data["business"]
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(*INK)
    pdf.cell(0, 9, _latin(b["name"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.set_text_color(*MUTED)
    details = " | ".join(x for x in (b["industry"], b.get("region"), f"Founded {b['founded']}",
                                     f"{b['employees']} employees") if x)
    pdf.cell(0, 5, _latin(f"Financial summary  |  {details}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 5, _latin(f"Data as of {data['asOf']}  |  Prepared {data['generated']}  |  Amounts in {cur}"),
             new_x="LMARGIN", new_y="NEXT")

    # ---- key figures
    pdf.ln(4)
    hs = data["health"]
    cards = [("Operating cash", m(data["operatingCash"])), ("Revenue, last 12 months", m(data["pnl"]["revenue"])),
             ("Net margin", f"{data['pnl']['netMargin'] * 100:.1f}%"),
             ("Health score", f"{hs['score']:.0f}/100 {hs['grade']}" if hs.get("score") is not None else "Not enough data")]
    cw = (W - 9) / 4
    y = pdf.get_y()
    for i, (label, value) in enumerate(cards):
        x = 15 + i * (cw + 3)
        pdf.set_fill_color(*SOFT)
        pdf.rect(x, y, cw, 18, style="F")
        pdf.set_xy(x + 3, y + 2.5)
        pdf.set_font("Helvetica", size=8)
        pdf.set_text_color(*MUTED)
        pdf.cell(cw - 6, 4, _latin(label))
        pdf.set_xy(x + 3, y + 8)
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(*INK)
        pdf.cell(cw - 6, 7, _latin(value))
    pdf.set_y(y + 20)

    # ---- profit and loss
    p = data["pnl"]
    heading(f"Profit and loss, {p['start']} to {p['end']}")
    lines = [("Revenue", p["revenue"]), ("Cost of goods sold", -p["cogs"]), ("Gross profit", p["grossProfit"]),
             ("Operating expenses", -p["operatingExpenses"]), ("Operating profit", p["operatingProfit"]),
             ("Interest (estimated)", -p["interest"]), ("Taxes", -p["taxes"]),
             ("Owner compensation", -p["ownerCompensation"]), ("Net profit", p["netProfit"])]
    top = pdf.get_y()
    table(["Line item", "Amount"], [[k, m(v)] for k, v in lines], [62, 34], bold_last=True)
    pdf.set_font("Helvetica", size=8)
    pdf.set_text_color(*MUTED)
    pdf.cell(96, 5, _latin(f"Gross margin {p['grossMargin'] * 100:.1f}%  |  Net margin {p['netMargin'] * 100:.1f}%"))
    pdf.set_text_color(*INK)
    # monthly revenue bars, to the right of the table
    months = data["monthly"]
    cx, cy, ch, cwid = 118, top + 4, 46, W - 103
    pdf.set_font("Helvetica", "B", 8.5)
    pdf.set_xy(cx, top)
    pdf.cell(cwid, 4, "Monthly revenue")
    peak = max([r["revenue"] for r in months] + [1])
    bw = cwid / max(1, len(months))
    for i, r in enumerate(months):
        h = ch * r["revenue"] / peak
        pdf.set_fill_color(*((180, 165, 255) if r["partial"] else (106, 53, 232)))
        pdf.rect(cx + i * bw + 0.6, cy + ch - h + 2, bw - 1.2, h, style="F")
        pdf.set_font("Helvetica", size=5.5)
        pdf.set_text_color(*MUTED)
        pdf.set_xy(cx + i * bw, cy + ch + 2.5)
        pdf.cell(bw, 3, r["month"][2:], align="C")
    pdf.set_text_color(*INK)
    pdf.set_y(max(pdf.get_y() + 6, cy + ch + 8))

    # ---- health score
    if hs.get("components"):
        heading(f"Financial health: {hs['score']:.0f}/100 ({hs['grade']}), weakest area: {hs['weakestArea']}")
        table(["Component", "Score", "Weight", "Why"],
              [[c["label"], f"{c['score']:.0f}", f"{c['weight'] * 100:.0f}%", c.get("explanation") or ""]
               for c in hs["components"]], [36, 14, 16, W - 66], text_cols=(0, 3))

    # ---- cash flow (page 2)
    pdf.add_page()
    f = data["forecast"]["summary"]
    heading("Cash position and 13-week forecast")
    facts = [["Opening balance (operating cash)", m(f["openingBalance"])],
             ["Lowest projected balance", f"{m(f['lowestBalance'])} in week {f['lowestWeek']}"],
             ["Safety buffer (about one week of outgoings)", m(f["safetyBuffer"])],
             ["First week below the buffer", str(f["firstBelowBufferWeek"]) if f.get("firstBelowBufferWeek") else "None"],
             ["Shortfall to cover", m(f["shortfallAmount"])], ["Balance after 13 weeks", m(f["endingBalance"])]]
    table(["Measure", "Value"], facts, [100, W - 100])
    # closing-balance line with the buffer
    weeks = data["forecast"]["weeks"]
    if weeks:
        pdf.ln(4)
        gx, gy, gw, gh = 25, pdf.get_y() + 2, W - 12, 40
        values = [w["closing"] for w in weeks] + [f["safetyBuffer"], 0]
        lo, hi = min(values), max(values)
        span = (hi - lo) or 1
        ypos = lambda v: gy + gh - (v - lo) / span * gh  # noqa: E731
        pdf.set_draw_color(*LINE)
        pdf.line(gx, ypos(0), gx + gw, ypos(0))
        pdf.set_draw_color(232, 163, 61)
        pdf.set_dash_pattern(dash=1.5, gap=1.2)
        pdf.line(gx, ypos(f["safetyBuffer"]), gx + gw, ypos(f["safetyBuffer"]))
        pdf.set_dash_pattern()
        pdf.set_draw_color(42, 120, 214)
        pdf.set_line_width(0.6)
        step = gw / max(1, len(weeks) - 1)
        for i in range(len(weeks) - 1):
            pdf.line(gx + i * step, ypos(weeks[i]["closing"]), gx + (i + 1) * step, ypos(weeks[i + 1]["closing"]))
        pdf.set_line_width(0.2)
        low = min(range(len(weeks)), key=lambda i: weeks[i]["closing"])
        pdf.set_fill_color(229, 72, 77)
        pdf.ellipse(gx + low * step - 1.2, ypos(weeks[low]["closing"]) - 1.2, 2.4, 2.4, style="F")
        pdf.set_font("Helvetica", size=7)
        pdf.set_text_color(*MUTED)
        for i, w in enumerate(weeks):
            if i % 2 == 0:
                pdf.set_xy(gx + i * step - 6, gy + gh + 2)
                pdf.cell(12, 3, w["start"][5:], align="C")
        pdf.set_xy(gx, gy + gh + 6)
        pdf.cell(gw, 4, _latin("Projected closing balance by week (blue), safety buffer (orange dashes), low point (red)."))
        pdf.set_text_color(*INK)
        pdf.set_y(gy + gh + 12)

    # ---- receivables
    r = data["receivables"]
    heading(f"Receivables: {m(r['totals']['total'])} open, average collection {r['dso']:.0f} days")
    buckets = ("current", "1-30", "31-60", "61-90", "90+")
    table(["Age", *("Current" if k == "current" else f"{k} days" for k in buckets), "Total"],
          [["All customers", *(m(r["totals"][k]) for k in buckets), m(r["totals"]["total"])]],
          [30] + [(W - 54) / 5] * 5 + [24])
    overdue = [c for c in r["customers"] if c["oldestDaysPastDue"] > 0][:5]
    if overdue:
        pdf.ln(2)
        table(["Most overdue customers", "Open invoices", "Oldest (days)", "Amount owed"],
              [[c["customer"], str(c["invoices"]), str(c["oldestDaysPastDue"]), m(c["total"])] for c in overdue],
              [W - 90, 30, 30, 30])

    # ---- loans
    heading("Existing debt")
    loans = data["loans"]
    if loans:
        table(["Lender", "Type", "Outstanding", "Rate", "Monthly payment"],
              [[x["lender"], x["type"], m(x["outstanding"]), f"{x['rate'] * 100:.1f}%", m(x["monthlyPayment"])]
               for x in loans], [W - 120, 34, 30, 22, 34], text_cols=(0, 1))
    else:
        pdf.set_font("Helvetica", size=9)
        pdf.cell(0, 6, "No existing loans recorded.", new_x="LMARGIN", new_y="NEXT")

    # ---- method note
    heading("How these figures were produced")
    pdf.set_font("Helvetica", size=8.5)
    pdf.set_text_color(*MUTED)
    pdf.multi_cell(0, 4.4, _latin(
        "All figures are calculated by deterministic, tested finance engines from the transactions, invoices and bills "
        "this business imported or recorded; no figure is produced by an AI model. Revenue is money received for sales; "
        "interest is estimated from recorded loan balances and rates. The forecast combines open invoices (adjusted for "
        "each customer's usual payment delay), bills, detected recurring payments and seasonal sales. These statements "
        "are unaudited management figures prepared for discussion, not financial advice."))
    return bytes(pdf.output())
