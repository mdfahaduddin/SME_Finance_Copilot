"""Reports: figures agree with the analysis engine, filters work, and CSV / Excel / PDF downloads are safe."""
import csv
import io
from datetime import date, timedelta

import pdfplumber
import pytest
from openpyxl import load_workbook

from app import models
from app.engines import reports
from app.engines.analysis import profit_and_loss as trailing_pnl
from app.engines.analysis import receivables as receivables_summary


# ---- engine ------------------------------------------------------------------------------------

def test_pnl_matches_the_analysis_engine_for_the_same_window(ledger):
    """The report and the dashboard/copilot must never disagree: same 365 days, same figures."""
    end = ledger.as_of.date()
    report = reports.profit_and_loss(ledger, end - timedelta(days=364), end)
    expected = trailing_pnl(ledger, 365)
    total = {r["label"]: r["values"][-1] for r in report["rows"] if r["values"] is not None}
    assert total["Total revenue"] == pytest.approx(expected["revenue"], abs=0.05)
    assert total["Total cost of goods sold"] == pytest.approx(expected["cogs"], abs=0.05)
    assert total["Total operating expenses"] == pytest.approx(expected["operatingExpenses"], abs=0.05)
    assert total["Interest (estimated from loan balances)"] == pytest.approx(expected["interest"], abs=0.5)
    assert total["Net profit"] == pytest.approx(expected["netProfit"], abs=0.5)


def test_pnl_monthly_columns_add_up(ledger):
    start, end = reports.default_range(ledger)
    report = reports.profit_and_loss(ledger, start, end)
    assert report["columns"][-1] == "Total" and len(report["columns"]) == 13  # 12 calendar months + total
    for row in report["rows"]:
        if row["values"] is not None:
            assert sum(row["values"][:-1]) == pytest.approx(row["values"][-1], abs=0.05), row["label"]
    lines = {r["label"]: r["values"][-1] for r in report["rows"] if r["values"] is not None}
    assert lines["Gross profit"] == pytest.approx(lines["Total revenue"] - lines["Total cost of goods sold"], abs=0.05)
    assert report["partialMonths"] == [report["columns"][-2]]  # the current month is not complete yet


def test_invalid_periods_are_refused(ledger):
    with pytest.raises(reports.ReportError):
        reports.profit_and_loss(ledger, date(2026, 5, 1), date(2026, 4, 1))
    with pytest.raises(reports.ReportError):
        reports.profit_and_loss(ledger, date(2020, 1, 1), date(2026, 1, 1))


def test_aged_receivables_reconcile(ledger):
    report = reports.aged_receivables(ledger)
    summary = receivables_summary(ledger)
    assert report["totals"]["total"] == pytest.approx(summary["openTotal"], abs=0.05)
    for bucket, amount in summary["aging"].items():
        assert report["totals"][bucket] == pytest.approx(amount, abs=0.05)
    first = report["customers"][0]
    assert first["customer"] == "Riverside Hotel Group" and first["oldestDaysPastDue"] > 30  # most overdue first


# ---- API ---------------------------------------------------------------------------------------

def test_reports_json_for_every_role(client):
    viewer = client.as_role("viewer")
    for kind in ("pnl", "receivables", "forecast", "transactions"):
        r = client.get(f"/api/reports/{kind}", headers=viewer)
        assert r.status_code == 200, (kind, r.text)
        assert r.json()["currency"] == "USD"
    opts = client.get("/api/reports/options", headers=viewer).json()
    assert opts["accounts"] and all(a["number"].startswith("****") for a in opts["accounts"]) and opts["categories"]
    assert client.get("/api/reports/pnl").status_code == 401  # signed-in users only


def test_transaction_filters(client):
    h = client.as_role("owner")
    base = client.get("/api/reports/transactions", headers=h).json()
    rent = client.get("/api/reports/transactions", params={"category": "Rent"}, headers=h).json()
    assert 0 < rent["count"] < base["count"] and {r["category"] for r in rent["rows"]} == {"Rent"}
    incoming = client.get("/api/reports/transactions", params={"direction": "in"}, headers=h).json()
    assert incoming["totals"]["out"] == 0 and all(r["signed"] > 0 for r in incoming["rows"])
    account = client.get("/api/reports/options", headers=h).json()["accounts"][0]
    one = client.get("/api/reports/transactions", params={"accountId": account["id"]}, headers=h).json()
    assert one["count"] and all(r["account"].startswith(account["name"]) for r in one["rows"])
    found = client.get("/api/reports/transactions", params={"q": "riverside"}, headers=h).json()
    assert found["count"] and all("riverside" in (r["description"] + r["counterparty"]).lower() for r in found["rows"])
    narrow = client.get("/api/reports/transactions", params={"start": "2026-09-01", "end": "2026-09-30"}, headers=h).json()
    assert narrow["count"] and all("2026-09-01" <= r["date"] <= "2026-09-30" for r in narrow["rows"])
    assert client.get("/api/reports/transactions", params={"start": "2026-09-30", "end": "2026-09-01"},
                      headers=h).status_code == 422


def test_csv_and_excel_downloads(client):
    h = client.as_role("accountant")
    for kind in ("pnl", "receivables", "forecast", "transactions"):
        c = client.get(f"/api/reports/{kind}", params={"format": "csv"}, headers=h)
        assert c.status_code == 200 and c.headers["content-type"].startswith("text/csv")
        assert c.headers["content-disposition"].startswith('attachment; filename="lotus-bakery-cafe-ltd-')
        rows = list(csv.reader(io.StringIO(c.content.decode("utf-8-sig"))))
        assert len(rows) > 1 and "(USD)" in ",".join(rows[0])
        x = client.get(f"/api/reports/{kind}", params={"format": "xlsx"}, headers=h)
        assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
        ws = load_workbook(io.BytesIO(x.content)).active
        assert ws["A1"].value and "Lotus Bakery" in ws["A2"].value
    # The P&L sheet's net-profit total equals the JSON report.
    pnl = client.get("/api/reports/pnl", headers=h).json()
    net = next(r for r in pnl["rows"] if r["label"] == "Net profit")["values"][-1]
    ws = load_workbook(io.BytesIO(client.get("/api/reports/pnl?format=xlsx", headers=h).content)).active
    row = next(r for r in ws.iter_rows(values_only=True) if r[0] == "Net profit")
    assert row[-1] == pytest.approx(net)


def test_spreadsheet_formula_injection_is_neutralised(client):
    """A statement description like '=HYPERLINK(...)' must be text in the export, never a formula."""
    h = client.as_role("owner")
    with client.session() as s:
        account = s.query(models.SmeBankAccount).filter_by(BusinessId=1).first()
        s.add(models.SmeTransaction(AccountId=account.Id, TxnDate=date(2026, 10, 1), Amount=10, Direction="out",
                                    Category="Other Expense", Counterparty="@evil",
                                    Description='=HYPERLINK("http://evil.example","click")'))
        s.commit()
    params = {"q": "evil", "format": "csv"}
    text = client.get("/api/reports/transactions", params=params, headers=h).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'@evil" in text
    params["format"] = "xlsx"
    ws = load_workbook(io.BytesIO(client.get("/api/reports/transactions", params=params, headers=h).content)).active
    cell = next(c for r in ws.iter_rows() for c in r if isinstance(c.value, str) and "HYPERLINK" in c.value)
    assert cell.data_type == "s"  # a string, not a formula


def test_lender_summary_pdf(client):
    h = client.as_role("owner")
    r = client.get("/api/reports/summary.pdf", headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content.startswith(b"%PDF")
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        assert len(pdf.pages) == 2
    for expected in ("Lotus Bakery & Cafe Ltd.", "Financial summary", "Profit and loss", "Net profit",
                     "13-week forecast", "Receivables", "Existing debt", "not financial advice", "Riverside Hotel Group"):
        assert expected in text, expected
    actions = [a["action"] for a in client.get("/api/audit?limit=20", headers=h).json()["items"]]
    assert "report.downloaded" in actions


def test_reports_are_isolated_and_in_the_business_currency(client):
    r = client.post("/api/auth/register", json={
        "businessName": "Green Leaf Grocers", "industry": "Retail", "region": "Dhaka", "foundedDate": "2021-04-01",
        "employees": 6, "currency": "BDT", "fullName": "Karim Uddin", "email": "karim@greenleaf.example",
        "password": "correct horse battery staple", "acceptTerms": True})
    assert r.status_code == 201, r.text
    tx = client.get("/api/reports/transactions").json()
    assert tx["currency"] == "BDT" and tx["count"] == 0  # none of the bakery's rows
    assert client.get("/api/reports/receivables").json()["customers"] == []
    head = client.get("/api/reports/pnl", params={"format": "csv"}).content.decode("utf-8-sig").splitlines()[0]
    assert "(BDT)" in head
    pdf = client.get("/api/reports/summary.pdf")
    assert pdf.status_code == 200  # an empty new business still gets a valid summary
    with pdfplumber.open(io.BytesIO(pdf.content)) as doc:
        text = "\n".join(p.extract_text() or "" for p in doc.pages)
    assert "Green Leaf Grocers" in text and "BDT" in text and "Lotus" not in text


def test_copilot_points_to_reports(client):
    h = client.as_role("owner")
    sid = client.post("/api/chat/sessions", headers=h).json()["id"]
    answer = client.post(f"/api/chat/sessions/{sid}/messages?stream=false",
                         json={"content": "Can I download a profit and loss report for my bank?"}, headers=h).json()
    assert "**Reports**" in answer["content"] and "PDF" in answer["content"]
    assert answer["groundedness"]["unverified"] == []
